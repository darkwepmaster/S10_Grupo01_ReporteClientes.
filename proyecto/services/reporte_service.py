from pathlib import Path
from datetime import datetime
import logging

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

logger = logging.getLogger(__name__)

COLUMNAS = ["ID", "Nombres", "Apellidos", "Edad"]


def construir_dataframe(registros):
    df = pd.DataFrame(registros)
    if df.empty:
        return df
    df.columns = COLUMNAS
    return df


def exportar_clientes(df: pd.DataFrame, carpeta: str = "reportes") -> Path:
    carpeta = Path(carpeta)
    carpeta.mkdir(exist_ok=True)

    nombre = f"clientes_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    ruta = carpeta / nombre

    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Clientes")
        ws = writer.book["Clientes"]

        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="D5002B")
            celda.alignment = Alignment(horizontal="center")

        for col in ws.columns:
            ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = ancho + 4

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    logger.info("Reporte exportado: %s (%s registros)", ruta, len(df))
    return ruta
